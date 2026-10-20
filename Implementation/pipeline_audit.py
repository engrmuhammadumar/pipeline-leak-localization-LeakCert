"""Read-only first-pass audit of pipeline datasets. Python 3.10+.

CLI: python pipeline_audit.py --root "F:\\Pipeline Data\\pipeline data"
All statistics describe a bounded HEAD sample, not a complete recording.
No labels, sensor roles, sample rate, experimental groups or chronology are inferred.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import math
import os
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from itertools import islice
from pathlib import Path

import numpy as np
import pandas as pd

TEXT_EXTENSIONS = {".csv", ".tsv", ".txt", ".dat"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
MISSING_TOKENS = {"", "nan", "na", "n/a", "null", "none", "nat"}
VERSION = "1.0.0"


def filename_hints(name):
    """Hypotheses only: never use these fields as ground-truth labels."""
    stem = Path(name).stem.lower()
    fluid = "gas" if re.search(r"(?:^|[_\W])gas(?:$|[_\W])", stem) else ""
    if re.search(r"(?:^|[_\W])water(?:$|[_\W])", stem):
        fluid = "water"
    match = re.search(r"(\d+(?:\.\d+)?)\s*bar", stem)
    return {"fluid_hint_unverified": fluid,
            "pressure_bar_hint_unverified": float(match.group(1)) if match else None}


def _encoding(path):
    with Path(path).open("rb") as f:
        sample = f.read(65536)
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    if b"\x00" in sample:
        raise ValueError("NUL bytes found: possibly binary or UTF-16 without a BOM. Specify encoding.")
    for enc in ("utf-8-sig", "cp949", "cp1252", "latin1"):
        try:
            # A prefix can end mid-character; keep an incomplete final character buffered.
            import codecs
            codecs.getincrementaldecoder(enc)().decode(sample, final=False)
            return enc
        except UnicodeError:
            continue
    raise ValueError("Cannot determine a text encoding; provide encoding explicitly.")


def _frame(rows, *, max_rows, max_cols, header_row, skiprows, info):
    """Keep the first row unless a human explicitly supplies header_row."""
    if header_row is not None and (not isinstance(header_row, int) or header_row < 0):
        raise ValueError("header_row must be None or a nonnegative integer after skiprows.")
    offset = 0 if header_row is None else header_row + 1
    sampled, observed_widths = [], []
    for row in islice(rows, skiprows, skiprows + offset + max_rows + 1):
        observed_widths.append(len(row))
        sampled.append(list(row[:max_cols + 1]))
    data = sampled[offset:offset + max_rows]
    widths = observed_widths[offset:offset + max_rows]
    width = min(max_cols, max(widths + ([len(sampled[header_row])] if
                    header_row is not None and len(sampled) > header_row else [0])))
    if header_row is None:
        names = [f"c{i:03d}" for i in range(width)]
    else:
        if len(sampled) <= header_row:
            raise ValueError("The specified header row is beyond the available rows.")
        names, used = [], set()
        for i in range(width):
            v = sampled[header_row][i] if i < len(sampled[header_row]) else None
            base = str(v).strip() if v is not None and str(v).strip() else f"c{i:03d}"
            name, k = base, 1
            while name in used:
                name, k = f"{base}__{k}", k + 1
            names.append(name)
            used.add(name)
    normalized = [(list(r[:width]) + [None] * width)[:width] for r in data]
    df = pd.DataFrame(normalized, columns=names)
    info.update({"sample_rows": len(df), "sample_columns": len(df.columns),
                 "more_rows_after_sample": len(sampled) > offset + max_rows,
                 "columns_truncated": any(w > max_cols for w in observed_widths),
                 "raw_width_min": min(widths, default=0),
                 "raw_width_max": max(widths, default=0),
                 "uneven_row_widths": len(set(widths)) > 1,
                 "header_row": header_row, "skiprows": skiprows,
                 "sample_strategy": "head", "first_row_preserved": header_row is None})
    return df, info


def load_preview(path, *, sheet=None, max_rows=5000, max_cols=128,
                 header_row=None, skiprows=0, delimiter=None, encoding=None):
    """Read a bounded preview. Returns (DataFrame, parser_information).

    Raw audit defaults: header_row=None, generic c000/c001 column names.
    For a confirmed header: header_row=0. skiprows counts parsed records.
    delimiter='whitespace' supports simple whitespace-delimited numeric files.
    Excel previews use stored formula values; formulas are not recalculated.
    """
    path = Path(path)
    if max_rows < 1 or max_cols < 1 or skiprows < 0:
        raise ValueError("max_rows/max_cols must be positive; skiprows cannot be negative.")
    ext = path.suffix.lower()
    info = {"sheet": "", "encoding": "", "delimiter": "", "parser_notes": []}
    kwargs = dict(max_rows=max_rows, max_cols=max_cols, header_row=header_row,
                  skiprows=skiprows, info=info)
    if ext in TEXT_EXTENSIONS:
        enc = encoding or _encoding(path)
        with path.open("r", encoding=enc, errors="strict", newline="") as f:
            sample = f.read(65536)
            f.seek(0)
            sep = delimiter
            if sep is None:
                try:
                    sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    if ext == ".tsv":
                        sep = "\t"
                    elif ext == ".csv":
                        sep = ","
                    elif re.search(r"\S[ \t]+\S", sample):
                        sep = "whitespace"
                    else:
                        sep = ","  # also preserves a one-column file
                    info["parser_notes"].append("Delimiter fallback used; inspect the raw preview.")
            info.update(encoding=enc, delimiter=repr(sep))
            if encoding is None:
                info["parser_notes"].append("Encoding inferred; verify text and units.")
            if sep == "whitespace":
                info["parser_notes"].append("Whitespace splitting does not support quoted fields.")
                rows = (line.split() for line in f)
            else:
                if len(sep) != 1:
                    raise ValueError("delimiter must be one character or 'whitespace'.")
                rows = csv.reader(f, delimiter=sep, strict=True)
            return _frame(rows, **kwargs)
    if ext in {".xlsx", ".xlsm"}:
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = book.worksheets[0] if sheet is None else (book.worksheets[sheet]
                  if isinstance(sheet, int) else book[sheet])
            info.update(sheet=ws.title, worksheet_rows_hint=ws.max_row,
                        worksheet_columns_hint=ws.max_column)
            info["parser_notes"].append("Excel dimensions are hints; formula values are cached only.")
            # Read one extra column so truncation is detectable. Explicit row bound
            # avoids iterating through an accidentally formatted million-row sheet.
            limit = skiprows + (header_row + 1 if header_row is not None else 0) + max_rows + 1
            if ws.max_row is not None:
                limit = min(limit, ws.max_row)
            def rows():
                for row in ws.iter_rows(max_row=limit, max_col=max_cols + 1, values_only=True):
                    values = list(row)
                    while values and values[-1] is None:
                        values.pop()
                    yield values
            result = _frame(rows(), **kwargs)
            if ws.max_column and ws.max_column > max_cols:
                result[1]["columns_truncated"] = True
            return result
        finally:
            book.close()
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            ws = book.sheet_by_index(0 if sheet is None else sheet) if (
                 sheet is None or isinstance(sheet, int)) else book.sheet_by_name(sheet)
            info.update(sheet=ws.name, worksheet_rows_hint=ws.nrows,
                        worksheet_columns_hint=ws.ncols)
            info["parser_notes"].append("Legacy XLS reader can load workbook internals into memory.")
            return _frame((ws.row_values(i, 0, min(ws.ncols, max_cols + 1))
                           for i in range(ws.nrows)), **kwargs)
        finally:
            book.release_resources()
    raise ValueError(f"Preview is not implemented for {ext or '(no extension)'}. See inventory.")


def _sheets(path):
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()
    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return wb.sheet_names()
        finally:
            wb.release_resources()
    return [None]


def numeric_values(series, decimal="."):
    """Parse numeric-looking cells without assigning their scientific meaning."""
    if decimal not in {".", ","}:
        raise ValueError("decimal must be '.' or ','.")
    s = series.astype("string").str.strip()
    missing = s.isna() | s.str.lower().isin(MISSING_TOKENS)
    s = s.mask(missing)
    if decimal == ",":
        s = s.str.replace(",", ".", regex=False)
    numeric = pd.to_numeric(s, errors="coerce").astype(float)
    return numeric, missing


def profile_columns(df, table_id, decimal="."):
    rows = []
    for name in df.columns:
        x, missing = numeric_values(df[name], decimal=decimal)
        finite = x[np.isfinite(x)]
        count = int((~missing).sum())
        ratio = float(x.notna().sum() / count) if count else 0.0
        row = {"table_id": table_id, "column": name,
               "first_cell_unverified": str(df[name].iloc[0])[:120] if len(df) else "",
               "sample_rows": len(df), "missing_count": int(missing.sum()),
               "missing_fraction": float(missing.mean()) if len(df) else None,
               "numeric_fraction_of_nonmissing": ratio,
               "non_numeric_nonmissing_count": int((x.isna() & ~missing).sum()),
               "infinite_count": int(np.isinf(x).sum()),
               "unique_values_sample": int(df[name].nunique(dropna=True)),
               "numeric_candidate_unverified": ratio >= 0.95 and len(finite) >= 2,
               "finite_count": len(finite),
               "finite_min": float(finite.min()) if len(finite) else None,
               "finite_max": float(finite.max()) if len(finite) else None,
               "finite_mean": float(finite.mean()) if len(finite) else None,
               "finite_std": float(finite.std(ddof=1)) if len(finite) > 1 else None,
               "constant_finite_sample": bool(finite.nunique() == 1) if len(finite) else None}
        rows.append(row)
    return rows


def _plot_candidates(df, profiles, path, title, decimal):
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    candidates = [p["column"] for p in profiles if p["numeric_candidate_unverified"]][:4]
    if not candidates:
        return False
    fig = Figure(figsize=(10, 2.4 * len(candidates)), layout="constrained")
    FigureCanvasAgg(fig)
    axes = np.atleast_1d(fig.subplots(len(candidates), 1, squeeze=True))
    for ax, col in zip(axes, candidates):
        x, _ = numeric_values(df[col], decimal=decimal)
        ax.plot(np.arange(len(x)), x.where(np.isfinite(x)), lw=0.8, color="#16627b")
        ax.set(xlabel="Row offset in head sample (not confirmed time)", ylabel=col)
        ax.grid(alpha=0.2)
    fig.suptitle(title + "\nNumeric candidates only; confirm sensor roles and units", fontsize=10)
    fig.savefig(path, dpi=130)
    return True


def _clean_json(obj):
    if isinstance(obj, dict):
        return {str(k): _clean_json(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_clean_json(v) for v in obj]
    if isinstance(obj, (np.integer, np.floating)):
        obj = obj.item()
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    return obj


def _write_csv(rows, path, columns=None):
    pd.DataFrame(rows, columns=columns).to_csv(path, index=False, encoding="utf-8-sig")


def make_bundle(run_dir):
    """Recreate the shareable audit archive. Contains summaries and small previews."""
    run_dir = Path(run_dir)
    archive = run_dir / "audit_bundle.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        for path in sorted(run_dir.rglob("*")):
            if path.is_file() and path != archive:
                z.write(path, path.relative_to(run_dir).as_posix())
    return archive


def run_audit(root, out_dir=None, *, max_rows=5000, max_cols=128,
              plots=True, hash_files=False, overrides=None):
    """Audit all files without modifying raw data; each call creates a new run.

    overrides maps a relative filename to load_preview options plus decimal.
    Example: {'Gas_13Bar.csv': {'header_row': 0, 'delimiter': ';', 'decimal': ','}}.
    All Excel sheets use that file's overrides. Confirm unusual sheets individually
    with load_preview. Optional SHA-256 checks exact file duplicates, not shared runs.
    """
    root = Path(root).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Dataset folder not found: {root}\nRun this on the Windows PC containing your F: drive.")
    out_dir = Path(out_dir).expanduser().resolve() if out_dir else root.parent / "pipeline_audit_outputs"
    if out_dir == root or root in out_dir.parents:
        raise ValueError("Choose an output directory outside the raw dataset folder.")
    if max_rows < 1 or max_cols < 1:
        raise ValueError("Sample limits must be positive.")
    stamp = datetime.now(timezone.utc).strftime("run_%Y%m%dT%H%M%S_%fZ")
    run_dir = out_dir / stamp
    run_dir.mkdir(parents=True, exist_ok=False)
    (run_dir / "previews").mkdir()
    (run_dir / "plots").mkdir()
    inventory, tables, profiles, issues, metadata = [], [], [], [], []
    overrides = overrides or {}

    def problem(file, sheet, severity, message):
        issues.append(dict(file=file, sheet=sheet or "", severity=severity, message=message))

    def walk_error(error):
        problem("[directory traversal]", "", "error", str(error))

    paths = []
    for current, dirs, files in os.walk(root, followlinks=False, onerror=walk_error):
        dirs.sort()
        for name in sorted(files):
            paths.append(Path(current) / name)
    for index, path in enumerate(paths, 1):
        rel = path.relative_to(root).as_posix()
        print(f"[{index}/{len(paths)}] {rel}", flush=True)
        item = dict(file=rel, extension=path.suffix.lower(), status="pending", bytes=None,
                    sha256="", tables_previewed=0, **filename_hints(path.name))
        inventory.append(item)
        if path.is_symlink():
            item["status"] = "skipped_symlink"
            continue
        try:
            item["bytes"] = path.stat().st_size
            if hash_files:
                h = hashlib.sha256()
                with path.open("rb") as f:
                    for chunk in iter(lambda: f.read(1024 * 1024), b""):
                        h.update(chunk)
                item["sha256"] = h.hexdigest()
            ext = path.suffix.lower()
            if path.name.startswith("~$"):
                item["status"] = "skipped_excel_lock"
                continue
            if ext == ".mat":
                from scipy.io import whosmat
                item["mat_variables"] = [dict(name=n, shape=list(shape), matlab_class=k)
                                          for n, shape, k in whosmat(path)]
                item["status"] = "metadata_only"
                problem(rel, "", "info", "MAT variables inventoried; no signal data loaded. v7.3 needs a separate HDF5 reader.")
                continue
            if ext not in TEXT_EXTENSIONS | EXCEL_EXTENSIONS:
                item["status"] = "unsupported_format"
                problem(rel, "", "info", "File inventoried but its content was not parsed.")
                continue
            sheets = _sheets(path)
            item["sheet_count"] = len(sheets) if ext in EXCEL_EXTENSIONS else None
            failures = 0
            for sheet in sheets:
                try:
                    opts = dict(overrides.get(rel, {}))
                    decimal = opts.pop("decimal", ".")
                    opts.pop("sheet", None)
                    df, info = load_preview(path, sheet=sheet, max_rows=max_rows,
                                            max_cols=max_cols, **opts)
                    tid = hashlib.sha256(f"{rel}::{sheet}".encode()).hexdigest()[:16]
                    preview_rel = f"previews/{tid}.csv"
                    df.head(12).to_csv(run_dir / preview_rel, index=False, encoding="utf-8-sig")
                    profile = profile_columns(df, tid, decimal=decimal)
                    profiles.extend(profile)
                    record = dict(table_id=tid, file=rel, **info, decimal=decimal,
                                  duplicate_rows_in_sample=int(df.duplicated().sum()),
                                  preview_file=preview_rel, plot_file="")
                    tables.append(record)
                    item["tables_previewed"] += 1
                    if info["columns_truncated"]:
                        problem(rel, sheet, "warning", "Column limit reached; increase max_cols and verify table orientation.")
                    if info["uneven_row_widths"]:
                        problem(rel, sheet, "warning", "Uneven row widths; inspect delimiter, header, metadata and missing trailing fields.")
                    if not len(df):
                        problem(rel, sheet, "warning", "No rows were sampled.")
                    metadata.append(dict(table_id=tid, file=rel, sheet=sheet or "",
                        recording_id="", experiment_id="", row_start_inclusive="",
                        row_stop_exclusive="", confirmed_fluid="", confirmed_pressure_bar="",
                        sensor_columns="", sensor_units="", time_column="", time_units="",
                        sampling_rate_hz="", label_column="", label_mapping="",
                        fixed_label_if_whole_recording="", leak_size_value="", leak_size_units="",
                        leak_position_m="", pipe_length_m="", sensor_positions_m="",
                        verified_chronological_order="", source_and_collection_protocol="", notes=""))
                    if plots:
                        try:
                            plot_rel = f"plots/{tid}.png"
                            if _plot_candidates(df, profile, run_dir / plot_rel,
                                                rel + (f" / {sheet}" if sheet else ""), decimal):
                                record["plot_file"] = plot_rel
                        except Exception as error:
                            problem(rel, sheet, "warning", f"Plot failed: {type(error).__name__}: {error}")
                except Exception as error:
                    failures += 1
                    problem(rel, sheet, "error", f"{type(error).__name__}: {error}")
            item["status"] = "previewed" if failures == 0 else ("partial_error" if item["tables_previewed"] else "error")
        except Exception as error:
            item["status"] = "error"
            problem(rel, "", "error", f"{type(error).__name__}: {error}")
    for rel in set(overrides) - {x["file"] for x in inventory}:
        problem(rel, "", "warning", "Override did not match an inventoried relative filename.")
    hashes = Counter(i["sha256"] for i in inventory if i["sha256"])
    duplicates = [[i["file"] for i in inventory if i["sha256"] == key]
                  for key, count in hashes.items() if count > 1]
    versions = {}
    for package in ("numpy", "pandas", "matplotlib", "scipy", "openpyxl", "xlrd"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "not installed"
    report = dict(audit_version=VERSION, created_utc=stamp, dataset_folder_name=root.name,
                  config=dict(max_rows=max_rows, max_cols=max_cols, plots=plots,
                              hash_files=hash_files, overrides=overrides),
                  versions=versions, inventory=inventory, tables=tables, issues=issues,
                  exact_duplicate_file_groups=duplicates,
                  interpretation="Head samples only; no confirmed labels, sampling rate, run boundaries, chronology or results.")
    (run_dir / "report.json").write_text(json.dumps(_clean_json(report), indent=2,
                                                    ensure_ascii=False, allow_nan=False), encoding="utf-8")
    _write_csv(inventory, run_dir / "inventory.csv", None if inventory else ["file", "status"])
    _write_csv(tables, run_dir / "tables.csv", None if tables else ["table_id", "file", "sheet", "sample_rows"])
    _write_csv(profiles, run_dir / "column_profile.csv", None if profiles else ["table_id", "column"])
    _write_csv(issues, run_dir / "issues.csv", ["file", "sheet", "severity", "message"])
    _write_csv(metadata, run_dir / "recording_metadata_template.csv", None if metadata else ["table_id", "file", "experiment_id"])
    status = Counter(i["status"] for i in inventory)
    summary = ["# Pipeline dataset audit", "", f"Files inventoried: {len(inventory)}",
               f"Tables previewed: {len(tables)}", f"File status counts: {dict(status)}",
               f"Maximum sampled rows per table: {max_rows}",
               f"Exact duplicate file groups: {duplicates if hash_files else 'not checked'}", "",
               "All statistics use the first sampled rows. They are not full-dataset statistics.",
               "The default reader preserves a possible header as data, using c000/c001 column names.",
               "Missing-value counts include blank, NA, N/A, nan, null, None and NaT tokens.",
               "Review issues.csv, the small previews and the numeric candidate plots before interpreting anything.",
               "A head sample can miss every leak event. No fault absence or class balance is inferred.",
               "Do not use numeric-candidate columns automatically as model inputs.", "",
               "## Next information needed", "",
               "1. Confirm which rows/columns are time samples, sensors, labels, metadata or separate trials.",
               "2. Confirm sampling rate, time and physical units, sensor positions, pipe geometry and label meanings.",
               "3. Complete recording_metadata_template.csv from the collection protocol; never guess labels from filenames.",
               "4. A table can contain several experiments; add one metadata row per experiment segment with row bounds.",
               "5. Split independent experiments before windowing. Keep duplicates and related recordings in one partition.",
               "6. Fit preprocessing on training data only. Reserve target-test labels for final scoring.", "",
               "Send audit_bundle.zip and one representative original data file for dataset-specific analysis.",
               "The archive contains small data previews and summaries; it does not upload anything by itself.", ""]
    (run_dir / "report.md").write_text("\n".join(summary), encoding="utf-8")
    archive = make_bundle(run_dir)
    print(f"\nAudit saved to: {run_dir}\nShare this report: {archive}", flush=True)
    return run_dir


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--rows", type=int, default=5000)
    parser.add_argument("--cols", type=int, default=128)
    parser.add_argument("--no-plots", action="store_true")
    parser.add_argument("--hash-files", action="store_true", help="Read full files for SHA-256 duplicate checking.")
    parser.add_argument("--overrides", type=Path, help="Optional JSON mapping relative filenames to parser options.")
    args = parser.parse_args()
    overrides = json.loads(args.overrides.read_text(encoding="utf-8-sig")) if args.overrides else None
    run_audit(args.root, args.out, max_rows=args.rows, max_cols=args.cols,
              plots=not args.no_plots, hash_files=args.hash_files, overrides=overrides)


if __name__ == "__main__":
    main()
